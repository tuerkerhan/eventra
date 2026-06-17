import io
import os
import re
import zipfile
from xml.sax.saxutils import escape as xml_escape
from datetime import datetime

from docx import Document
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from sqlalchemy.orm import Session, joinedload

from ..core.deps import get_current_user
from ..database import get_db
from ..models import ContractTemplate, Event, EventFormFieldDef, EventType, EventTypeFieldDef, Salon, SalonUser
from ..schemas import ContractTemplateOut

router = APIRouter(prefix="/contracts", tags=["contracts"])

UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "uploads")

TURKISH_MONTHS = [
    "Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran",
    "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık",
]


def _fmt_date(s: str) -> str:
    if not s:
        return ""
    try:
        d = datetime.strptime(s, "%Y-%m-%d")
        return f"{d.day} {TURKISH_MONTHS[d.month - 1]} {d.year}"
    except Exception:
        return s


def _fmt_money(v: float) -> str:
    return f"₺{v:,.0f}".replace(",", ".")


def _build_replacements(event: Event, salon: Salon, custom_field_defs: list | None = None) -> dict[str, str]:
    event_type_name = event.event_type.name if event.event_type else ""
    randevu_no = f"{salon.contract_prefix}-{event.appointment_no}" if event.appointment_no else ""
    remaining = event.total_fee - event.total_paid
    gelin_damat = event.bride_groom or next(
        (cf.value for cf in (event.custom_fields or []) if cf.key == "gelin_damat"), ""
    )

    replacements = {
        "%baslik%": event.title or "",
        "%isim%": event.full_name or "",
        "%gelin_damat%": gelin_damat or "",
        "%tc_no%": event.tc_no or "",
        "%telefon%": event.mobile_phone or event.phone or "",
        "%email%": event.email or "",
        "%tarih%": _fmt_date(event.event_date),
        "%sozlesme_tarihi%": _fmt_date(event.contract_date) if event.contract_date else "",
        "%baslama_saati%": event.start_time or "",
        "%baslangic_saati%": event.start_time or "",
        "%bitis_saati%": event.end_time or "",
        "%tip%": event_type_name,
        "%etkinlik_turu%": event_type_name,
        "%rezervasyon_durumu%": event.reservation_status or "",
        "%davetli_sayisi%": str(event.guest_count),
        "%misafir_sayisi%": str(event.guest_count),
        "%adres%": event.address or "",
        "%bolge%": event.region or "",
        "%toplam_ucret%": _fmt_money(event.total_fee),
        "%kapora%": _fmt_money(event.kapora_amount),
        "%odenen%": _fmt_money(event.total_paid),
        "%kalan%": _fmt_money(remaining),
        "%randevu_no%": randevu_no,
        "%personel%": event.staff or "",
        "%notlar%": event.note or "—",
        "%not%": event.note or "—",
        "%salon_adi%": salon.name or "",
        "%salon_adresi%": salon.address or "",
        "%kdv_orani%": f"%{salon.vat_rate:.0f}",
    }

    # Merge custom field values. Custom fields can define an explicit contract
    # placeholder; otherwise the field key works as "%field_key%".
    if custom_field_defs:
        tag_map = {d.key: getattr(d, "placeholder_tag", "") or f"%{d.key}%" for d in custom_field_defs}
        for cf in (event.custom_fields or []):
            tag = tag_map.get(cf.key)
            if tag:
                replacements[tag] = cf.value or ""

    return replacements


# ─── .docx replacement ────────────────────────────────────────────────────────

def _replace_paragraph(para, replacements: dict[str, str]) -> None:
    """Merge all runs in a paragraph, replace placeholders, rewrite into first run."""
    if not para.runs:
        return
    full = "".join(r.text for r in para.runs)
    replaced = full
    for k, v in replacements.items():
        replaced = replaced.replace(k, v)
    if replaced == full:
        return
    para.runs[0].text = replaced
    for run in para.runs[1:]:
        run.text = ""


def _fill_docx(path: str, replacements: dict[str, str]) -> bytes:
    doc = Document(path)

    for para in doc.paragraphs:
        _replace_paragraph(para, replacements)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_paragraph(para, replacements)

    # Headers & footers
    for section in doc.sections:
        for hdr in (section.header, section.footer, section.even_page_header,
                    section.even_page_footer, section.first_page_header, section.first_page_footer):
            if hdr is not None:
                for para in hdr.paragraphs:
                    _replace_paragraph(para, replacements)
                for table in hdr.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for para in cell.paragraphs:
                                _replace_paragraph(para, replacements)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─── .odt replacement ────────────────────────────────────────────────────────

# LibreOffice splits placeholders across text:span elements in 3 known patterns.
# We only touch text:span tags to avoid accidentally consuming paragraph/heading tags.

# Pattern C: <text:span ...>%</text:span>  →  bare %
_PAT_C = re.compile(r'<text:span[^>]*>%</text:span>')
# Pattern A: %<text:span ...>key%</text:span ...>  →  %key%
_PAT_A = re.compile(r'%(?:<text:span[^>]*>)+([a-z_]+%)(?:</text:span>)*')
# Pattern B: %<text:span ...>key</text:span>%  →  %key%
_PAT_B = re.compile(r'%(?:<text:span[^>]*>)+([a-z_]+)(?:</text:span>)+%')


def _fix_split_placeholders(xml_str: str) -> str:
    xml_str = _PAT_C.sub('%', xml_str)       # isolated % spans → bare %
    xml_str = _PAT_A.sub(r'%\1', xml_str)    # %<span>key%</span> → %key%
    xml_str = _PAT_B.sub(r'%\1%', xml_str)   # %<span>key</span>% → %key%
    return xml_str


def _fill_odt(path: str, replacements: dict[str, str]) -> bytes:
    """Replace placeholders in ODT, handling text split across XML span elements."""
    buf = io.BytesIO()
    with zipfile.ZipFile(path, "r") as zin:
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename in ("content.xml", "styles.xml"):
                    text = data.decode("utf-8")
                    # Collapse LibreOffice-split placeholders first
                    text = _fix_split_placeholders(text)
                    # Now do straightforward replacement (XML-escape the values)
                    for k, v in replacements.items():
                        text = text.replace(k, xml_escape(v))
                    data = text.encode("utf-8")
                zout.writestr(item, data)
    return buf.getvalue()


# ─── Routes ───────────────────────────────────────────────────────────────────

@router.get("/templates", response_model=list[ContractTemplateOut])
def list_templates(db: Session = Depends(get_db), user: SalonUser = Depends(get_current_user)):
    return (
        db.query(ContractTemplate)
        .filter(ContractTemplate.salon_id == user.salon_id)
        .order_by(ContractTemplate.created_at.desc())
        .all()
    )


@router.post("/templates/upload", response_model=ContractTemplateOut)
async def upload_template(
    file: UploadFile = File(...),
    name: str = "",
    db: Session = Depends(get_db),
    user: SalonUser = Depends(get_current_user),
):
    if not file.filename:
        raise HTTPException(status_code=400, detail="Dosya adı boş")

    ext = file.filename.rsplit(".", 1)[-1].lower()
    if ext not in ("docx", "odt"):
        raise HTTPException(status_code=400, detail="Sadece .docx ve .odt desteklenir")

    salon_dir = os.path.join(UPLOAD_DIR, str(user.salon_id))
    os.makedirs(salon_dir, exist_ok=True)

    tmpl = ContractTemplate(
        salon_id=user.salon_id,
        name=name.strip() or file.filename.rsplit(".", 1)[0],
        original_filename=file.filename,
        file_type=ext,
        file_path="",  # set after we know the id
    )
    db.add(tmpl)
    db.flush()

    file_path = os.path.join(salon_dir, f"{tmpl.id}.{ext}")
    content = await file.read()
    with open(file_path, "wb") as f:
        f.write(content)

    tmpl.file_path = file_path
    db.commit()
    db.refresh(tmpl)
    return tmpl


@router.delete("/templates/{template_id}", status_code=204)
def delete_template(
    template_id: str,
    db: Session = Depends(get_db),
    user: SalonUser = Depends(get_current_user),
):
    tmpl = db.query(ContractTemplate).filter(
        ContractTemplate.id == template_id,
        ContractTemplate.salon_id == user.salon_id,
    ).first()
    if not tmpl:
        raise HTTPException(status_code=404, detail="Şablon bulunamadı")
    if os.path.exists(tmpl.file_path):
        os.remove(tmpl.file_path)
    db.delete(tmpl)
    db.commit()


@router.get("/generate/{event_id}")
def generate_contract(
    event_id: str,
    template_id: str = Query(...),
    db: Session = Depends(get_db),
    user: SalonUser = Depends(get_current_user),
):
    event = (
        db.query(Event)
        .options(joinedload(Event.event_type), joinedload(Event.custom_fields))
        .filter(Event.id == event_id, Event.salon_id == user.salon_id)
        .first()
    )
    if not event:
        raise HTTPException(status_code=404, detail="Etkinlik bulunamadı")

    tmpl = db.query(ContractTemplate).filter(
        ContractTemplate.id == template_id,
        ContractTemplate.salon_id == user.salon_id,
    ).first()
    if not tmpl:
        raise HTTPException(status_code=404, detail="Şablon bulunamadı")

    if not os.path.exists(tmpl.file_path):
        raise HTTPException(status_code=404, detail="Şablon dosyası sunucuda bulunamadı")

    salon = db.query(Salon).filter(Salon.id == user.salon_id).first()
    custom_defs = db.query(EventFormFieldDef).filter(
        EventFormFieldDef.salon_id == user.salon_id,
        EventFormFieldDef.is_builtin == False,
    ).all()
    type_custom_defs = db.query(EventTypeFieldDef).filter(
        EventTypeFieldDef.salon_id == user.salon_id,
    ).all()
    custom_defs = [*custom_defs, *type_custom_defs]
    replacements = _build_replacements(event, salon, custom_defs)

    if tmpl.file_type == "docx":
        output = _fill_docx(tmpl.file_path, replacements)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        output = _fill_odt(tmpl.file_path, replacements)
        media_type = "application/vnd.oasis.opendocument.text"

    safe_title = re.sub(r"[^\w\s-]", "", event.title or event.full_name or "sozlesme").strip().replace(" ", "_")
    appointment_part = f"{salon.contract_prefix}-{event.appointment_no}" if event.appointment_no else event.id[:8]
    filename = f"sozlesme_{appointment_part}_{safe_title}_{event.event_date}.{tmpl.file_type}"

    return Response(
        content=output,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
